"""
PPTX parser using python-pptx.
Extracts:
  - Text per slide (title + body)
  - Images per slide → captioned by VLM
  - Speaker notes
"""
from __future__ import annotations
import asyncio
import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)


@dataclass
class SlideContent:
    slide_num: int           # 1-indexed
    title: str
    body_text: str
    notes: str
    image_captions: list[str] = field(default_factory=list)
    raw_images: list[bytes] = field(default_factory=list)


async def parse_pptx(path: Path, caption_images: bool = True) -> list[SlideContent]:
    """Parse a PPTX file and return per-slide content."""
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image

    from app.services.llm.provider import get_vlm

    prs = Presentation(str(path))
    slides: list[SlideContent] = []
    vlm = get_vlm() if caption_images else None

    for slide_idx, slide in enumerate(prs.slides):
        slide_num = slide_idx + 1
        title = ""
        body_parts: list[str] = []
        raw_images: list[bytes] = []

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        if shape.shape_type == 13 or (hasattr(shape, "placeholder_format") and
                                                       shape.placeholder_format and
                                                       shape.placeholder_format.idx == 0):
                            title = text
                        else:
                            body_parts.append(text)

            # Extract images
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                try:
                    img_bytes = shape.image.blob
                    if len(img_bytes) > 5000:
                        # Convert to JPEG bytes for the VLM
                        img = Image.open(io.BytesIO(img_bytes))
                        buf = io.BytesIO()
                        img.convert("RGB").save(buf, format="JPEG", quality=85)
                        raw_images.append(buf.getvalue())
                except Exception as e:
                    logger.warning(f"Failed to extract image from slide {slide_num}: {e}")

        # Speaker notes
        notes = ""
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes = "\n".join(p.text for p in notes_tf.paragraphs if p.text.strip())

        # Caption images
        captions: list[str] = []
        if raw_images and vlm:
            caption_tasks = [
                _caption_slide_image(vlm, img, slide_num, i)
                for i, img in enumerate(raw_images)
            ]
            captions = await asyncio.gather(*caption_tasks)

        slides.append(SlideContent(
            slide_num=slide_num,
            title=title,
            body_text="\n".join(body_parts),
            notes=notes,
            image_captions=captions,
            raw_images=raw_images,
        ))
        logger.debug(f"Slide {slide_num}: title={repr(title)}, {len(raw_images)} images")

    return slides


async def _caption_slide_image(vlm, img_bytes: bytes, slide_num: int, img_idx: int) -> str:
    from tenacity import retry, stop_after_attempt, wait_exponential

    @retry(stop=stop_after_attempt(3), wait=wait_exponential(multiplier=1, min=2, max=10))
    async def _call():
        response = await vlm.generate(
            prompt=(
                f"You are analyzing an image from slide {slide_num} of a lecture presentation. "
                "Describe what is shown: if it's a diagram, explain its structure and key concepts; "
                "if it's a graph/chart, describe axes, trends, and values; "
                "if it's code, transcribe and explain it; "
                "if it's a photo or illustration, describe it in a way a student would find useful. "
                "Be concise but thorough (2-5 sentences)."
            ),
            images=[img_bytes],
            max_tokens=512,
            temperature=0.1,
        )
        return response.text.strip()

    try:
        return await _call()
    except Exception as e:
        logger.error(f"Slide image captioning failed (slide {slide_num}, img {img_idx}): {e}")
        return f"[Image on slide {slide_num} — captioning failed]"
